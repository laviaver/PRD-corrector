"""
File parser utility for extracting text content from various file formats.

All formats are converted to markdown (.md) format before analysis.
Supports .txt, .md, .docx - all converted to markdown.
Conversion output is trimmed to reduce irrelevant words and shorten the file.
"""

import re
from typing import BinaryIO

import mammoth
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from src.models.prd import PRDFileType
from src.utils.logger import get_logger

logger = get_logger(__name__)

# Boilerplate phrases to drop from converted markdown (case-insensitive)
_TRIM_DROP_LINES = frozenset({
    "page break", "pagebreak", "[page break]", "[pagebreak]",
    "confidential", "draft", "internal use only", "proprietary",
    "table of contents", "continued", "continued from previous page",
    "---", "***", "…", "",
})
# Drop "Page X" / "Page X of Y" style lines
_TRIM_PAGE_OF_RE = re.compile(r"^page\s+\d+(\s+of\s+\d+)?$", re.IGNORECASE)

# Markdown image: ![alt](url) — url can be long (e.g. data:image/...;base64,...)
_IMAGE_MD_RE = re.compile(r"!\[([^\]]*)\]\([^)]+\)", re.DOTALL)


def _replace_images_with_placeholder(content: str) -> str:
    """
    Replace markdown images with a short note so the LLM knows an image was there
    without embedding large base64 or file content. Reduces token count.
    """
    def repl(match: re.Match) -> str:
        alt = (match.group(1) or "").strip()
        if alt:
            return f"[Image: {alt}]"
        return "[Image]"

    return _IMAGE_MD_RE.sub(repl, content)


def _trim_markdown(content: str) -> str:
    """
    Shorten markdown: replace images with placeholders, collapse newlines, strip lines,
    drop boilerplate, page-only lines, and punctuation-only lines; collapse spaces.
    """
    if not content or not content.strip():
        return content.strip()
    content = _replace_images_with_placeholder(content)
    lines = content.splitlines()
    out = []
    prev_empty = False
    for line in lines:
        s = line.strip()
        s = re.sub(r"[ \t]+", " ", s)
        if not s:
            if not prev_empty:
                out.append("")
            prev_empty = True
            continue
        lower = s.lower()
        if lower in _TRIM_DROP_LINES:
            continue
        if _TRIM_PAGE_OF_RE.match(lower):
            continue
        # Drop lines that are only a number (e.g. page number)
        if re.match(r"^\d+$", s):
            continue
        if len(s) <= 2 and re.match(r"^[\s\-_*.]*$", s):
            continue
        if re.match(r"^[\-\_*.\s]+$", s) and len(s) > 3:
            continue
        out.append(s)
        prev_empty = False
    text = "\n".join(out)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

# Style map so Word Heading 1/2/3 (and common titles) become markdown # / ## / ###
MAMMOTH_HEADING_STYLE_MAP = """
p[style-name='Heading 1'] => h1:fresh
p[style-name='Heading 2'] => h2:fresh
p[style-name='Heading 3'] => h3:fresh
p[style-name^='Heading'] => h2:fresh
p[style-name='Title'] => h1:fresh
p[style-name='Section Title'] => h2:fresh
p[style-name='Subsection Title'] => h3:fresh
"""

# Word style names that map to markdown heading levels (for python-docx fallback)
DOCX_HEADING_PREFIX = {
    "heading 1": "# ",
    "heading 2": "## ",
    "heading 3": "### ",
    "title": "# ",
}


def _docx_fallback_with_headings(tmp_path: str, filename: str) -> str:
    """
    Fallback when mammoth returns empty: use python-docx and prepend # / ## / ###
    for paragraphs with Heading 1/2/3 or Title style so structure extractor can split.
    """
    logger.warning(f"Mammoth extracted no content from {filename}, trying python-docx fallback")
    doc = Document(tmp_path)
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "").lower()
        prefix = DOCX_HEADING_PREFIX.get(style_name)
        if prefix:
            parts.append(prefix + text)
        else:
            parts.append(text)
    content = "\n".join(parts)
    if len(content) < 100 and hasattr(doc, "tables") and doc.tables:
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        table_texts.append(cell.text.strip())
        if table_texts:
            content = "\n".join(parts + [""] + table_texts) if parts else "\n".join(table_texts)
    return content.strip()


class FileParseError(Exception):
    """Exception raised when file parsing fails."""

    def __init__(self, message: str, file_type: str = None):
        self.message = message
        self.file_type = file_type
        super().__init__(self.message)


def parse_file(file_content: BinaryIO, filename: str, file_type: PRDFileType) -> str:
    """
    Parse file content and convert to markdown format.

    Args:
        file_content: Binary file content
        filename: Original filename
        file_type: Type of file (txt, md, docx)

    Returns:
        Markdown-formatted text content

    Raises:
        FileParseError: If parsing fails
    """
    try:
        if file_type == PRDFileType.MD:
            # Already markdown, just read and return
            raw = file_content.read()
            for encoding in ('utf-8', 'cp1252', 'latin-1'):
                try:
                    content = raw.decode(encoding)
                    if encoding != 'utf-8':
                        logger.warning(
                            "File %s is not UTF-8; decoded with %s",
                            filename,
                            encoding,
                        )
                    return _trim_markdown(content.strip())
                except UnicodeDecodeError:
                    continue
            raise FileParseError("Failed to decode file content", file_type.value)

        elif file_type == PRDFileType.TXT:
            # Convert .txt to markdown format
            raw = file_content.read()
            for encoding in ('utf-8', 'cp1252', 'latin-1'):
                try:
                    content = raw.decode(encoding)
                    if encoding != 'utf-8':
                        logger.warning(
                            "File %s is not UTF-8; decoded with %s",
                            filename,
                            encoding,
                        )
                    # Convert plain text to markdown by preserving structure
                    # Wrap in code block or format as markdown paragraphs
                    content = content.strip()
                    if not content:
                        raise FileParseError("File is empty", file_type.value)
                    # Format as markdown: preserve line breaks and structure
                    lines = content.split('\n')
                    markdown_lines = []
                    for line in lines:
                        line = line.strip()
                        if line:
                            markdown_lines.append(line)
                        else:
                            markdown_lines.append('')
                    return _trim_markdown('\n'.join(markdown_lines))
                except UnicodeDecodeError:
                    continue
            raise FileParseError("Failed to decode file content", file_type.value)

        elif file_type == PRDFileType.DOCX:
            # Convert .docx to markdown using mammoth
            try:
                import tempfile
                import os

                # Read file content
                file_content.seek(0)
                docx_bytes = file_content.read()

                # Save to temporary location for mammoth
                with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                    tmp_file.write(docx_bytes)
                    tmp_file.flush()
                    tmp_path = tmp_file.name

                try:
                    # Use mammoth to convert .docx to markdown with explicit heading style map
                    # so # / ## / ### appear for structure extraction
                    with open(tmp_path, 'rb') as docx_file:
                        result = mammoth.convert_to_markdown(
                            docx_file,
                            style_map=MAMMOTH_HEADING_STYLE_MAP,
                            include_default_style_map=True,
                        )
                        content = result.value

                        # Log any warnings from mammoth
                        if result.messages:
                            for message in result.messages:
                                logger.warning(f"Mammoth conversion warning: {message}")

                    content = content.strip()
                    if not content:
                        content = _docx_fallback_with_headings(tmp_path, filename)
                    if not content:
                        raise FileParseError(
                            "Document has no extractable text. The file may be empty, "
                            "or text may be in images/headers/footers.",
                            file_type.value,
                        )
                    return _trim_markdown(content)
                finally:
                    # Clean up temporary file
                    if os.path.exists(tmp_path):
                        os.unlink(tmp_path)

            except PackageNotFoundError:
                raise FileParseError(f"Invalid .docx file: {filename}", file_type.value)
            except FileParseError:
                raise
            except Exception as e:
                logger.error(f"Error parsing .docx file {filename}: {e}")
                raise FileParseError(f"Failed to parse .docx file: {str(e)}", file_type.value)

        else:
            # Only .docx, .md, .txt are supported
            raise FileParseError(
                f"Unsupported file type: {file_type.value}. Only .docx, .md, and .txt files are allowed.",
                file_type.value,
            )

    except FileParseError:
        raise  # already a clear message (e.g. empty docx)
    except UnicodeDecodeError as e:
        logger.error(f"Encoding error parsing file {filename}: {e}")
        raise FileParseError(f"Failed to decode file content: {str(e)}", file_type.value)
    except Exception as e:
        logger.error(f"Unexpected error parsing file {filename}: {e}")
        raise FileParseError(f"Unexpected error parsing file: {str(e)}", file_type.value)


def parse_text_content(content: str) -> str:
    """
    Parse direct text content input.

    Args:
        content: Text content string

    Returns:
        Stripped text content

    Raises:
        ValueError: If content is empty or whitespace-only
    """
    if not content or not content.strip():
        raise ValueError("Content cannot be empty")

    return content.strip()
